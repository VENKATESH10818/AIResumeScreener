"""
Resume parsing service.

Supports:
  - PDF  → pdfplumber (text-based) with pytesseract fallback for scanned pages
  - DOCX → python-docx
  - Images (PNG/JPG/TIFF/BMP/WEBP) → pytesseract

After text extraction, spaCy NER + regex patterns extract structured fields:
  name, email, phone, skills, job_titles, companies, education, certifications,
  total_years_experience (estimated).
"""

from __future__ import annotations

import hashlib
import io
import logging
import os
import platform
import re
import sys
from pathlib import Path
from typing import Any

import pdfplumber
import pytesseract
import spacy
from docx import Document
from PIL import Image

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Tesseract path — Windows requires an explicit path if not on system PATH
# ---------------------------------------------------------------------------

def _configure_tesseract() -> None:
    """
    Auto-detect the Tesseract executable on Windows.
    On Linux/macOS it is expected to be on PATH (installed via apt/brew).
    Can always be overridden with the TESSERACT_CMD env var.
    """
    custom_cmd = os.getenv("TESSERACT_CMD")
    if custom_cmd:
        pytesseract.pytesseract.tesseract_cmd = custom_cmd
        return

    if platform.system() == "Windows":
        candidates = [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ]
        for path in candidates:
            if Path(path).exists():
                pytesseract.pytesseract.tesseract_cmd = path
                logger.info("Tesseract found at: %s", path)
                return
        logger.warning(
            "Tesseract not found at default Windows paths. "
            "Install from https://github.com/UB-Mannheim/tesseract/wiki "
            "or set TESSERACT_CMD env var. OCR features will be unavailable."
        )


_configure_tesseract()


# ---------------------------------------------------------------------------
# Lazy-loaded spaCy model
# ---------------------------------------------------------------------------

_NLP: spacy.language.Language | None = None
SPACY_MODEL = os.getenv("SPACY_MODEL", "en_core_web_sm")


def _get_nlp() -> spacy.language.Language:
    global _NLP
    if _NLP is None:
        logger.info("Loading spaCy model: %s", SPACY_MODEL)
        _NLP = spacy.load(SPACY_MODEL)
    return _NLP


# ---------------------------------------------------------------------------
# Skill ontology (extend as needed — load from a JSON file for production)
# ---------------------------------------------------------------------------

SKILL_KEYWORDS: set[str] = {
    # Programming languages
    "python", "java", "javascript", "typescript", "c++", "c#", "go", "rust",
    "ruby", "php", "swift", "kotlin", "scala", "r", "matlab", "sql", "bash",
    # Frameworks / libraries
    "django", "flask", "fastapi", "spring", "react", "angular", "vue",
    "node.js", "express", "tensorflow", "pytorch", "keras", "scikit-learn",
    "pandas", "numpy", "spark", "hadoop", "kafka",
    # Cloud & DevOps
    "aws", "gcp", "azure", "docker", "kubernetes", "terraform", "ansible",
    "ci/cd", "github actions", "jenkins", "linux", "git",
    # Data / ML / AI
    "machine learning", "deep learning", "nlp", "computer vision", "llm",
    "data analysis", "data engineering", "etl", "nosql", "mongodb",
    "postgresql", "mysql", "redis", "elasticsearch", "faiss",
    # Soft / domain skills
    "agile", "scrum", "project management", "communication", "leadership",
    "rest api", "graphql", "microservices", "system design",
}

# ---------------------------------------------------------------------------
# Regex patterns
# ---------------------------------------------------------------------------

EMAIL_RE = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
PHONE_RE = re.compile(
    r"(\+?\d{1,3}[\s\-.]?)?(\(?\d{2,4}\)?[\s\-.]?)?\d{3,4}[\s\-.]?\d{4}"
)
EXP_YEARS_RE = re.compile(
    r"(\d+)\s*\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:experience|exp\.?|work)",
    re.IGNORECASE,
)
EDUCATION_KEYWORDS = {
    "bachelor", "b.sc", "b.s.", "b.tech",
    "master", "m.sc", "m.s.", "m.tech", "mba", "m.e.",
    "phd", "ph.d", "doctorate",
    "associate", "diploma", "certification",
}
CERT_KEYWORDS = {
    "aws certified", "google certified", "azure certified",
    "pmp", "cissp", "cpa", "cfa", "gcp professional",
    "tensorflow developer", "deep learning specialization",
}


# ---------------------------------------------------------------------------
# Text extraction helpers
# ---------------------------------------------------------------------------


def _extract_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF. Falls back to OCR on pages without selectable text."""
    texts: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if len(page_text.strip()) < 20:
                # Likely scanned — OCR the page image
                try:
                    img = page.to_image(resolution=200).original
                    page_text = pytesseract.image_to_string(img)
                except Exception as exc:
                    logger.warning("OCR fallback failed on PDF page: %s", exc)
            texts.append(page_text)
    return "\n".join(texts)


def _extract_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX preserving paragraph order."""
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)


def _extract_image(file_bytes: bytes) -> str:
    """OCR an image file (PNG / JPG / TIFF / BMP / WEBP)."""
    img = Image.open(io.BytesIO(file_bytes))
    return pytesseract.image_to_string(img)


def extract_text(file_bytes: bytes, file_ext: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = file_ext.lower().strip().lstrip(".")
    if ext == "pdf":
        return _extract_pdf(file_bytes)
    elif ext in {"docx", "doc"}:
        return _extract_docx(file_bytes)
    elif ext in {"png", "jpg", "jpeg", "tiff", "tif", "bmp", "webp"}:
        return _extract_image(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")


def detect_file_type(filename: str) -> str:
    """Return a normalised file type string."""
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "pdf":
        return "pdf"
    if ext in {"docx", "doc"}:
        return "docx"
    if ext in {"png", "jpg", "jpeg", "tiff", "tif", "bmp", "webp"}:
        return "image"
    return "unknown"


# ---------------------------------------------------------------------------
# NLP extraction helpers
# ---------------------------------------------------------------------------


def _extract_name(doc: spacy.tokens.Doc, raw_text: str) -> str | None:
    """
    Extract the candidate's name from a resume.

    Strategy (in order):
      1. Look for a PERSON entity in the first 30 lines only — names are
         always near the top; entities deeper in the doc are companies/places.
      2. Fall back to the first line that looks like a proper name (2-4 words,
         all capitalised, no digits, no special chars, not a known non-name).
    """

    # --- Tokens / patterns that are never a person's name ---
    _NAME_BLOCKLIST = {
        "gmail", "outlook", "yahoo", "hotmail", "icloud",
        "resume", "curriculum vitae", "cv",
        "profile", "contact", "contacts", "name", "address",
        "phone", "email", "mobile", "tel",
        "linkedin", "github", "gitlab", "portfolio", "website",
        "objective", "summary", "career objective",
        "microsoft", "google", "apple", "facebook", "amazon", "netflix",
        "java", "python", "javascript", "typescript", "swift", "kotlin",
        "c++", "c#", "ruby", "golang", "scala", "rust",
        "skills", "education", "experience", "work experience",
        "references", "projects", "achievements", "awards",
    }

    # Pattern: 2-4 words, each starting with capital letter, only letters/hyphens/spaces
    # Matches "Yashwanth Kumar", "Alice Mary Johnson", "Jean-Pierre Martin"
    _NAME_RE = re.compile(r"^[A-Z][a-zA-Z\-']{1,30}(?:\s+[A-Z][a-zA-Z\-']{1,30}){1,3}$")

    def _is_valid_name(text: str) -> bool:
        t = text.strip()
        if not t:
            return False
        if t.lower() in _NAME_BLOCKLIST:
            return False
        if any(c.isdigit() for c in t):
            return False
        if re.search(r"[/@#$%^&*()+=\[\]{}<>|\\:;\"!?]", t):
            return False
        if len(t.split()) < 2 or len(t.split()) > 5:
            return False
        return bool(_NAME_RE.match(t))

    # 1. Search only the first 30 lines of text (names are always at the top)
    first_30_lines = "\n".join(raw_text.splitlines()[:30])
    # Re-run spaCy on just the header section for better precision
    nlp = _get_nlp()
    header_doc = nlp(first_30_lines[:2000])

    for ent in header_doc.ents:
        if ent.label_ == "PERSON" and _is_valid_name(ent.text):
            return ent.text.strip()

    # 2. Line-by-line fallback — scan only first 20 lines
    for line in raw_text.splitlines()[:20]:
        line = line.strip()
        if (
            _is_valid_name(line)
            and not EMAIL_RE.search(line)
            and not PHONE_RE.search(line)
        ):
            return line

    return None


def _redact_name(name: str | None) -> str | None:
    """Return 'First L.' to minimise PII stored in DB."""
    if not name:
        return None
    parts = name.split()
    if len(parts) == 1:
        return parts[0]
    return f"{parts[0]} {parts[-1][0]}."


def _extract_emails(text: str) -> list[str]:
    return EMAIL_RE.findall(text)


def _extract_phones(text: str) -> list[str]:
    return [m.group().strip() for m in PHONE_RE.finditer(text)]


def _extract_skills(text: str) -> list[str]:
    """Case-insensitive keyword scan against SKILL_KEYWORDS."""
    lower = text.lower()
    return sorted({skill for skill in SKILL_KEYWORDS if skill in lower})


def _extract_job_titles(doc: spacy.tokens.Doc) -> list[str]:
    title_pattern = re.compile(
        r"\b(software engineer|data scientist|machine learning engineer|"
        r"backend developer|frontend developer|full[- ]?stack developer|"
        r"devops engineer|data analyst|product manager|engineering manager|"
        r"tech lead|architect|research scientist|ml engineer|ai engineer)\b",
        re.IGNORECASE,
    )
    titles = [m.group().strip().title() for m in title_pattern.finditer(doc.text)]
    return list(dict.fromkeys(titles))  # deduplicate, preserve order


def _extract_companies(doc: spacy.tokens.Doc) -> list[str]:
    return list({ent.text.strip() for ent in doc.ents if ent.label_ == "ORG"})


def _extract_education(text: str) -> list[str]:
    results: list[str] = []
    for line in text.splitlines():
        lower = line.lower()
        if any(kw in lower for kw in EDUCATION_KEYWORDS):
            cleaned = line.strip()
            if cleaned:
                results.append(cleaned)
    return results


def _extract_certifications(text: str) -> list[str]:
    lower = text.lower()
    return [cert for cert in CERT_KEYWORDS if cert in lower]


def _estimate_years_experience(text: str) -> float:
    """Return the maximum explicitly stated years-of-experience, or 0."""
    matches = EXP_YEARS_RE.findall(text)
    if not matches:
        return 0.0
    return float(max(int(y) for y in matches))


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


class ParsedResume:
    """Structured output of the parsing pipeline."""

    def __init__(self, raw_text: str, parsed: dict[str, Any]):
        self.raw_text = raw_text
        self.parsed = parsed

    def to_dict(self) -> dict[str, Any]:
        return self.parsed


def parse_resume(file_bytes: bytes, filename: str) -> ParsedResume:
    """
    Full pipeline: extract text → clean → NLP → return ParsedResume.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        filename:   Original filename (used to detect type).

    Returns:
        ParsedResume with .raw_text and .parsed dict.
    """
    file_ext = Path(filename).suffix.lstrip(".")  # e.g. "pdf", "docx"
    raw_text = extract_text(file_bytes, file_ext)

    # Basic cleaning
    clean_text = re.sub(r"\s{3,}", "\n", raw_text).strip()

    nlp = _get_nlp()
    doc = nlp(clean_text[:100_000])  # truncate to avoid OOM on huge docs

    name = _extract_name(doc, clean_text)
    emails = _extract_emails(clean_text)
    phones = _extract_phones(clean_text)
    skills = _extract_skills(clean_text)
    titles = _extract_job_titles(doc)
    companies = _extract_companies(doc)
    education = _extract_education(clean_text)
    certifications = _extract_certifications(clean_text)
    years_exp = _estimate_years_experience(clean_text)

    email_hash: str | None = None
    if emails:
        email_hash = hashlib.sha256(emails[0].lower().encode()).hexdigest()

    parsed: dict[str, Any] = {
        "name": name,
        "name_redacted": _redact_name(name),
        "email_hash": email_hash,
        "emails": emails,
        "phones": phones,
        "skills": skills,
        "job_titles": titles,
        "companies": companies,
        "education": education,
        "certifications": certifications,
        "years_experience": years_exp,
        "file_type": detect_file_type(filename),
        "char_count": len(clean_text),
    }

    return ParsedResume(raw_text=clean_text, parsed=parsed)


def parse_job_description(jd_text: str) -> dict[str, Any]:
    """
    Extract required/desired skills and metadata from a job description.
    """
    nlp = _get_nlp()
    doc = nlp(jd_text[:100_000])

    all_skills = _extract_skills(jd_text)
    titles = _extract_job_titles(doc)

    required_section_re = re.compile(
        r"(?:required|must[- ]have|mandatory|essential)[^\n]*\n(.*?)(?=\n\n|\Z)",
        re.IGNORECASE | re.DOTALL,
    )
    desired_section_re = re.compile(
        r"(?:nice[- ]to[- ]have|preferred|desired|bonus|optional)[^\n]*\n(.*?)(?=\n\n|\Z)",
        re.IGNORECASE | re.DOTALL,
    )

    req_text = " ".join(m.group(1) for m in required_section_re.finditer(jd_text))
    des_text = " ".join(m.group(1) for m in desired_section_re.finditer(jd_text))

    required_skills = _extract_skills(req_text) if req_text.strip() else all_skills
    desired_skills = _extract_skills(des_text) if des_text.strip() else []

    return {
        "all_skills": all_skills,
        "required_skills": required_skills,
        "desired_skills": desired_skills,
        "job_titles": titles,
        "years_experience_required": _estimate_years_experience(jd_text),
        "education": _extract_education(jd_text),
    }
